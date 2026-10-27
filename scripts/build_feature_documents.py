"""Build the two v2.0.1 delivery documents from reviewed Markdown.

Author with the bundled document Python. Render with the documents skill's
renderer and a controlled portable office runtime, never a running desktop
document session. Rendering is not a claim of visual acceptance.
"""
from __future__ import annotations
import argparse
from datetime import datetime, timezone
import hashlib
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
from urllib.parse import urljoin

ROOT = Path(__file__).resolve().parents[1]
SPECS = {
    'core': ('CORE_FEATURES_GUIDE', 'core_features_qa', '六项核心功能操作与验证'),
    'development': ('DEVELOPMENT_AND_DELIVERY', 'development_delivery_qa', '项目开发与交付过程'),
}
BASE = 'https://github.com/aluola/taintlab/blob/v2.0.1/docs/'

def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()

def blocks(text):
    lines = text.splitlines(); index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1; continue
        if line.startswith('```'):
            code = []; index += 1
            while index < len(lines) and not lines[index].strip().startswith('```'):
                code.append(lines[index]); index += 1
            if index == len(lines): raise ValueError('Unclosed code fence')
            yield 'code', '\n'.join(code); index += 1; continue
        if line.startswith('|'):
            table = []
            while index < len(lines) and lines[index].strip().startswith('|'):
                row = [x.strip() for x in lines[index].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?', x.replace(' ', '')) for x in row): table.append(row)
                index += 1
            if not table or any(len(r) != len(table[0]) for r in table): raise ValueError('Irregular Markdown table')
            yield 'table', table; continue
        if re.fullmatch(r'!\[[^\]]*\]\([^)]+\)', line):
            yield 'image', line; index += 1; continue
        heading = re.match(r'^(#{1,4})\s+(.+)', line)
        if heading:
            yield 'heading', (len(heading[1]), heading[2]); index += 1; continue
        bullet = re.match(r'^(?:[-*]|\d+[.、])\s+(.+)', line)
        if bullet:
            yield 'bullet', bullet[1]; index += 1; continue
        paragraph = [line]; index += 1
        while index < len(lines) and lines[index].strip() and not re.match(r'^(?:#|\||!\[|```|[-*]\s|\d+[.、]\s)', lines[index].strip()):
            paragraph.append(lines[index].strip()); index += 1
        yield 'paragraph', ' '.join(paragraph)

def make_architecture():
    from PIL import Image, ImageDraw, ImageFont
    target = ROOT/'docs/assets/development-delivery'; target.mkdir(parents=True, exist_ok=True)
    im = Image.new('RGB', (1680, 1000), '#ffffff'); d = ImageDraw.Draw(im)
    face = Path(os.environ.get('SystemRoot', 'C:/Windows'))/'Fonts/msyh.ttc'
    f = lambda size: ImageFont.truetype(str(face), size)
    d.text((55, 35), '架构与执行流程示意  运行核心 2.0.0', fill='#142f32', font=f(36))
    d.text((55, 89), '仅为代码结构示意  不代表模型内部因果或某次运行结果', fill='#596c6d', font=f(24))
    boxes = [
        (55,180,340,335,'用户与 React 页面','显式选择资料与目标','自然语言任务 / 快捷入口'),
        (380,180,665,335,'FastAPI 与 Service','Host Origin CSRF / 归属','服务端 TaskGrant / 事件'),
        (705,180,990,335,'Agent 执行循环','run_case / resume_run','预算 停止 反馈 / 动作提议'),
        (1030,180,1315,335,'确定性 Gateway','格式 白名单 敏感度','额度 一次性确认'),
        (1355,180,1630,335,'模拟 World','文档与本地快照','收件箱 / 虚拟文件'),
        (705,440,990,595,'本地 llama.cpp','Qwen3 0.6B Q8_0','只返回结构化动作提议'),
        (1030,440,1315,595,'Registry 来源 DAG','不可变标签与父节点','确定性派生 / 潜在依赖'),
        (55,700,540,870,'安全投影与同源 SSE','先检查可见性  后向浏览器展示','不发送未检查的原始模型 token'),
        (950,700,1630,870,'独立评测与真实记录回放','只读实际状态和预设断言  输出指标与证据','实验答案不进入 Agent 或 Gateway'),
    ]
    for x1,y1,x2,y2,a,b,c in boxes:
        d.rounded_rectangle((x1,y1,x2,y2),radius=14,fill='#f3f8f7',outline='#297e73',width=3)
        for j,text in enumerate((a,b,c)):
            d.text((x1+16,y1+20+43*j),text,fill='#173c3c',font=f(22 if j else 25))
    def arrow(start,end,label=None):
        import math
        d.line((start,end),fill='#457777',width=4); ang=math.atan2(end[1]-start[1],end[0]-start[0])
        tip=[end,(end[0]-14*math.cos(ang-.5),end[1]-14*math.sin(ang-.5)),(end[0]-14*math.cos(ang+.5),end[1]-14*math.sin(ang+.5))]
        d.polygon(tip,fill='#457777')
        if label:d.text(((start[0]+end[0])/2+12,(start[1]+end[1])/2-20),label,fill='#465d5c',font=f(21))
    for x in (340,665,990,1315): arrow((x,255),(x+40,255))
    arrow((825,335),(825,440))
    arrow((850,440),(850,335),'本机请求与提议')
    arrow((1170,440),(1170,335),'标签检查')
    arrow((1490,335),(1490,700),'保存实际状态')
    d.line(((1400,335),(1400,650),(600,650),(600,780)),fill='#457777',width=4)
    arrow((600,780),(540,780))
    d.text((750,610),'执行回执经安全投影',fill='#465d5c',font=f(21))
    arrow((180,700),(180,335),'获准事件')
    d.text((55,930),'模拟边界始终存在  none 也不能访问真实邮箱 数据库 Shell 或宿主任意路径',fill='#596c6d',font=f(25))
    out=target/'architecture.png'; im.save(out)
    (target/'architecture.json').write_text(json.dumps({'kind':'programmatic_design_schematic','runtime_version':'2.0.0','not_a_runtime_screenshot':True,'sha256':digest(out),'generated_by':'scripts/build_feature_documents.py --architecture-only'},ensure_ascii=False,indent=2),encoding='utf-8')

def build(kind):
    from docx import Document
    from docx.enum.section import WD_ORIENT, WD_SECTION
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
    from PIL import Image
    stem, qa_name, label = SPECS[kind]; source = ROOT/'docs'/f'{stem}.md'
    text = source.read_text('utf-8-sig')
    if '<!-- AUTO ' in text: raise ValueError('Actual evidence placeholders remain: '+str(source))
    doc=Document(); section=doc.sections[0]
    for border in list(doc.styles.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    def section_layout(landscape=False):
        section.orientation=WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
        section.page_width,section.page_height=(Cm(29.7),Cm(21)) if landscape else (Cm(21),Cm(29.7))
        section.top_margin=Cm(1.5);section.bottom_margin=Cm(1.4)
        section.left_margin=section.right_margin=Cm(1.65)
        section.header_distance=section.footer_distance=Cm(.65)
    section_layout()
    for name in ('Normal','Title','Heading 1','Heading 2','Heading 3','Caption','List Bullet'):
        style=doc.styles[name];style.font.name='Microsoft YaHei';style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
        style.font.color.rgb=RGBColor.from_string('142F32') if name not in ('Title','Normal') else RGBColor(0,0,0)
    doc.styles['Normal'].font.size=Pt(10.5);doc.styles['Normal'].paragraph_format.line_spacing=1.15
    doc.styles['Normal'].paragraph_format.space_after=Pt(6)
    doc.styles['Title'].font.size=Pt(26);doc.styles['Title'].paragraph_format.space_after=Pt(16)
    for name,size in [('Heading 1',18),('Heading 2',13),('Heading 3',11)]:
        doc.styles[name].font.size=Pt(size);doc.styles[name].paragraph_format.space_before=Pt(10);doc.styles[name].paragraph_format.space_after=Pt(6)
    doc.styles['Caption'].font.size=Pt(9)
    header=section.header.paragraphs[0];header.text=f'TaintLab  |  {label}  |  发布 v2.0.1  核心 2.0.0';header.style='Caption'
    footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('第 ');field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field);footer.add_run(' 页')
    def inline(p,value):
        for token in re.split(r'(\[[^\]]+\]\([^)]+\)|`[^`]+`|\*\*[^*]+\*\*)',value):
            m=re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)',token)
            if m:
                target=m[2] if '://' in m[2] else urljoin(BASE,m[2]);link=OxmlElement('w:hyperlink')
                link.set(qn('r:id'),p.part.relate_to(target,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True))
                r=OxmlElement('w:r');props=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),'226D68');props.append(c);r.append(props)
                t=OxmlElement('w:t');t.text=m[1];r.append(t);link.append(r);p._p.append(link)
            else:
                r=p.add_run(token[1:-1] if token.startswith('`') else token.replace('**',''));r.bold=token.startswith('**')
                if token.startswith('`'):r.font.name='Consolas';r.font.size=Pt(8.5)
    assets={};headings=[];compact_chapter=False;landscape_active=False
    def new_section(landscape):
        nonlocal section
        section=doc.add_section(WD_SECTION.NEW_PAGE);section_layout(landscape)
    for typ,value in blocks(text):
        if typ!='image' and landscape_active:
            new_section(False);landscape_active=False
        if typ=='heading':
            level,title=value
            if level==2:compact_chapter=kind=='development' and title.startswith('5 ')
            if level==1:doc.add_paragraph(title,'Title')
            else:
                if level==2:headings.append(title)
                heading=doc.add_paragraph(title,'Heading '+str(min(level-1,3)))
                if level==2:heading.paragraph_format.page_break_before=True
                elif compact_chapter:heading.paragraph_format.space_before=Pt(8)
        elif typ=='code':
            p=doc.add_paragraph();p.paragraph_format.line_spacing=1.05;p.paragraph_format.space_before=Pt(5);p.paragraph_format.keep_together=True
            r=p.add_run(value);r.font.name='Consolas';r.font.size=Pt(8.5)
        elif typ=='image':
            m=re.fullmatch(r'!\[([^\]]*)\]\(([^)]+)\)',value);path=(source.parent/m[2]).resolve()
            if not path.is_relative_to(ROOT) or not path.is_file():raise ValueError('Missing/outside image '+str(path))
            with Image.open(path) as im:w,h=im.size
            landscape=w>1000 and h>500
            if landscape:
                new_section(True);landscape_active=True
            elif landscape_active:
                new_section(False);landscape_active=False
            p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
            scale=min((26.3 if landscape else 17.6)/w,(15.5 if landscape else 18.0)/h)
            p.add_run().add_picture(str(path),width=Cm(w*scale))
            caption=doc.add_paragraph(m[1],'Caption');caption.alignment=WD_ALIGN_PARAGRAPH.CENTER
            assets[path.relative_to(ROOT).as_posix()]=digest(path)
        elif typ=='table':
            table=doc.add_table(rows=1,cols=len(value[0]));table.autofit=False
            count=len(value[0]);widths=[17.6/count]*count
            if count==2:widths=[6.0,11.6]
            elif count==3:widths=[4.2,6.7,6.7]
            elif count==6:widths=[1.0,2.3,4.1,5.1,2.1,3.0]
            if compact_chapter:
                for column,width in zip(table.columns,widths):column.width=Cm(width)
            for i,row in enumerate(value):
                cells=table.rows[0].cells if i==0 else table.add_row().cells
                trpr=table.rows[i]._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
                if i==0:trpr.append(OxmlElement('w:tblHeader'))
                for j,cell in enumerate(cells):
                    cell.width=Cm(widths[j]);p=cell.paragraphs[0];inline(p,row[j]);p.paragraph_format.space_after=Pt(3 if compact_chapter else 5);p.paragraph_format.space_before=Pt(3 if compact_chapter else 5)
                    for run in p.runs:run.font.size=Pt(9 if count<5 else 8.5);run.bold=i==0
                    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'DEEEEA' if i==0 else ('F5F8F7' if i%2==0 else 'FFFFFF'));cell._tc.get_or_add_tcPr().append(shade)
            doc.add_paragraph().paragraph_format.space_after=Pt(1)
        else:
            p=doc.add_paragraph(style='List Bullet' if typ=='bullet' else None);inline(p,value)
            if compact_chapter:p.paragraph_format.space_after=Pt(4)
    doc.core_properties.title=text.splitlines()[0].lstrip('# ');doc.core_properties.author='TaintLab 项目组（成员信息待补充）'
    doc.core_properties.subject='发布 v2.0.1；冻结运行核心 2.0.0；全部业务模拟'
    out=source.with_suffix('.docx');doc.save(out)
    qa=ROOT/'docs'/qa_name;qa.mkdir(parents=True,exist_ok=True)
    (qa/'build.json').write_text(json.dumps({'release_version':'2.0.1','runtime_version':'2.0.0','built_at_utc':datetime.now(timezone.utc).isoformat(),'source_sha256':digest(source),'docx_sha256':digest(out),'chapters':headings,'assets':assets,'layout':'A4, dedicated landscape pages for dense full UI screenshots; real crops retain original content'},ensure_ascii=False,indent=2),encoding='utf-8')
    return out,qa

def render(docx,qa):
    renderer=Path(os.environ.get('TAINTLAB_DOC_RENDERER',Path.home()/'.codex/plugins/cache/openai-primary-runtime/documents/26.1020.11348/skills/documents/render_docx.py'))
    runtime=Path(os.environ.get('TAINTLAB_DOC_RUNTIME',Path.home()/'.cache/codex-runtimes/codex-primary-runtime/dependencies'))
    office=Path(os.environ.get('TAINTLAB_SOFFICE',ROOT/'docs/qa/tools/LibreOffice/program/soffice.exe')).resolve()
    if not renderer.is_file() or not office.is_file():raise RuntimeError('Set TAINTLAB_DOC_RENDERER and TAINTLAB_SOFFICE to controlled document tools; do not use desktop session')
    output=qa/'render';output.mkdir(parents=True,exist_ok=True)
    env=dict(os.environ,PATH=os.pathsep.join([str(office.parent),str(runtime/'native/poppler/Library/bin'),os.environ.get('PATH','')]))
    result=subprocess.run([sys.executable,str(renderer),str(docx),'--output_dir',str(output),'--emit_pdf','--width','1800'],capture_output=True,env=env,encoding='utf-8',errors='replace')
    (qa/'render.log').write_text(result.stdout+'\n'+result.stderr,encoding='utf-8')
    if result.returncode:raise RuntimeError('Render failed; inspect '+str(qa/'render.log'))
    pdf=output/docx.with_suffix('.pdf').name
    if not pdf.exists():raise RuntimeError('No rendered PDF')
    shutil.copyfile(pdf,docx.with_suffix('.pdf'))
    from pypdf import PdfReader
    pages=len(PdfReader(pdf).pages)
    for path in output.glob('page-*.png'):
        m=re.fullmatch(r'page-(\d+)\.png',path.name)
        if m and int(m[1])>pages:
            old=qa/'prior_render'/datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%S')/path.name
            old.parent.mkdir(parents=True,exist_ok=True);path.replace(old)
    print(json.dumps({'document':docx.name,'pages':pages,'visual_review':'required_after_latest_render'}))

def main():
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('--document',choices=['core','development','all'],default='all');p.add_argument('--render',action='store_true');p.add_argument('--architecture-only',action='store_true');a=p.parse_args()
    if a.architecture_only:make_architecture();return
    for kind in (SPECS if a.document=='all' else [a.document]):
        document,qa=build(kind)
        if a.render:render(document,qa)

if __name__=='__main__':main()
