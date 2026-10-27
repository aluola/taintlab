"""Build the public manual from Markdown and real annotated captures.

Requires python-docx, Pillow and pypdf. PDF rendering uses the documents skill
renderer when TAINTLAB_DOC_RENDERER is set, or a locally installed LibreOffice.
No network or model calls; the project report is never overwritten.
"""
from __future__ import annotations
import argparse
import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urljoin
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'docs/USER_MANUAL.md'
QA = ROOT / 'docs/publication/manual_qa'
BASE = 'https://github.com/aluola/taintlab/blob/v2.0.0/docs/'

def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()

def inline(paragraph, value):
    for token in re.split(r'(\[[^\]]+\]\([^)]+\)|`[^`]+`|\*\*[^*]+\*\*)', value):
        match = re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)', token)
        if match:
            link = OxmlElement('w:hyperlink')
            target = match[2] if '://' in match[2] else urljoin(BASE, match[2])
            rid = paragraph.part.relate_to(target, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
            link.set(qn('r:id'), rid)
            run = OxmlElement('w:r'); props = OxmlElement('w:rPr')
            color = OxmlElement('w:color'); color.set(qn('w:val'), '245080'); props.append(color)
            run.append(props); text = OxmlElement('w:t'); text.text = match[1]; run.append(text); link.append(run)
            paragraph._p.append(link)
        else:
            run = paragraph.add_run(token.strip('`') if token.startswith('`') else token.replace('**', ''))
            if token.startswith('**'): run.bold = True
            if token.startswith('`'): run.font.name = 'Consolas'; run.font.size = Pt(9.5)

def build():
    source = SOURCE.read_text(encoding='utf-8-sig')
    document = Document()
    # The installed python-docx template contains decorative blue paragraph
    # rules. Remove them rather than inheriting an unintended title treatment.
    for border in list(document.styles.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    section = document.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(1.25), Cm(1.2)
    section.left_margin = section.right_margin = Cm(1.55)
    section.header_distance = section.footer_distance = Cm(.55)
    for name in ('Normal', 'Title', 'Heading 1', 'Heading 2', 'Caption'):
        style = document.styles[name]
        style.font.name = 'Microsoft YaHei'
        style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
        style.font.color.rgb = RGBColor(0, 0, 0)
    normal = document.styles['Normal']
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08
    document.styles['Title'].font.size = Pt(27)
    document.styles['Heading 1'].font.size = Pt(17)
    document.styles['Heading 1'].paragraph_format.space_before = Pt(0)
    document.styles['Heading 1'].paragraph_format.space_after = Pt(6)
    document.styles['Caption'].font.size = Pt(9)
    header = section.header.paragraphs[0]
    header.text = 'TaintLab  |  操作与实现手册  |  2.0 · v2.0.0'
    header.style = 'Caption'
    footer = section.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('本地模拟教学与研究  ·  第 ')
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
    footer.add_run(' 页')
    images, headings = [], []
    blocks = re.split(r'\n\s*\n', source.strip())
    def new_section(landscape):
        nonlocal section
        section = document.add_section(WD_SECTION.NEW_PAGE)
        section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
        section.page_width, section.page_height = (Cm(29.7), Cm(21)) if landscape else (Cm(21), Cm(29.7))
    dedicated = False
    for index, block in enumerate(blocks):
        if block.startswith('# '):
            document.add_paragraph(block[2:], 'Title')
        elif block.startswith('```'):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.keep_together = True
            paragraph.paragraph_format.line_spacing = 1.0
            run = paragraph.add_run('\n'.join(block.splitlines()[1:-1]))
            run.font.name = 'Consolas'; run.font.size = Pt(8.5)
        elif block.startswith('## '):
            title = block[3:]
            dedicated = False
            if index + 1 < len(blocks):
                upcoming = re.fullmatch(r'!\[([^\]]*)\]\(([^)]+)\)', blocks[index + 1])
                if upcoming:
                    with Image.open(SOURCE.parent / upcoming[2]) as picture:
                        dedicated = picture.width > 900 and picture.height / picture.width > .62
            new_section(dedicated)
            paragraph = document.add_paragraph(title, 'Heading 1')
            headings.append(title)
        elif re.fullmatch(r'!\[([^\]]*)\]\(([^)]+)\)', block):
            match = re.fullmatch(r'!\[([^\]]*)\]\(([^)]+)\)', block)
            path = SOURCE.parent / match[2]
            if not path.is_file(): raise FileNotFoundError(path)
            with Image.open(path) as picture: width, height = picture.size
            # Dense wide views get a dedicated, larger figure page. Sidebar
            # crops remain readable beside the ordinary page explanation.
            scale = min((26.3 if dedicated else 17.8) / width, (15.3 if dedicated else 10.5) / height)
            paragraph = document.add_paragraph(); paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.keep_with_next = True
            paragraph.add_run().add_picture(str(path), width=Cm(width * scale))
            caption = document.add_paragraph(match[1], 'Caption'); caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            images.append({'path': path.relative_to(ROOT).as_posix(), 'sha256': sha(path), 'pixels': [width, height]})
            if dedicated:
                new_section(False)
                continued = document.add_paragraph(headings[-1] + ' · 操作说明', 'Heading 1')
        else:
            inline(document.add_paragraph(), block.replace('\n', ' '))
    output = ROOT / 'docs/USER_MANUAL.docx'
    document.save(output)
    QA.mkdir(parents=True, exist_ok=True)
    (QA / 'build.json').write_text(json.dumps({'project_version': '2.0.0', 'built_at_utc': datetime.now(timezone.utc).isoformat(), 'source_markdown_sha256': sha(SOURCE), 'docx_sha256': sha(output), 'headings': headings, 'images': images, 'layout': 'A4 portrait explanations; dense screenshots on landscape pages; black headings; page fields'}, ensure_ascii=False, indent=2), encoding='utf-8')
    return output

def render(docx):
    runtime = Path(os.environ.get('TAINTLAB_DOC_RUNTIME', str(Path.home() / '.cache/codex-runtimes/codex-primary-runtime/dependencies')))
    renderer = Path(os.environ.get('TAINTLAB_DOC_RENDERER', str(Path.home() / '.codex/plugins/cache/openai-primary-runtime/documents/26.1020.11348/skills/documents/render_docx.py')))
    target = QA / 'render'; target.mkdir(parents=True, exist_ok=True)
    soffice = os.environ.get('TAINTLAB_SOFFICE') or shutil.which('soffice')
    if not soffice:
        candidate = ROOT / 'docs/qa/tools/LibreOffice/program/soffice.exe'
        if candidate.exists(): soffice = str(candidate)
    env = dict(os.environ)
    extra = [str(runtime / 'native/poppler/Library/bin')]
    if soffice: extra.insert(0, str(Path(soffice).parent))
    env['PATH'] = os.pathsep.join(extra + [env.get('PATH', '')])
    if renderer.is_file():
        command = [sys.executable, str(renderer), str(docx), '--output_dir', str(target), '--emit_pdf', '--width', '1800']
    elif soffice:
        command = [soffice, '--headless', '--convert-to', 'pdf', '--outdir', str(target), str(docx)]
    else: raise RuntimeError('Install LibreOffice or set TAINTLAB_DOC_RENDERER / TAINTLAB_SOFFICE; source DOCX is preserved.')
    completed = subprocess.run(command, env=env, capture_output=True, text=True, encoding='utf-8', errors='replace')
    (QA / 'render.log').write_text(completed.stdout + '\n' + completed.stderr, encoding='utf-8')
    if completed.returncode: raise RuntimeError('Rendering failed; see docs/publication/manual_qa/render.log')
    pdf = target / 'USER_MANUAL.pdf'
    if not pdf.exists(): raise RuntimeError('Renderer did not produce PDF')
    shutil.copyfile(pdf, ROOT / 'docs/USER_MANUAL.pdf')
    from pypdf import PdfReader
    pages = len(PdfReader(pdf).pages)
    for previous in target.glob('page-*.png'):
        number = re.fullmatch(r'page-(\d+)\.png', previous.name)
        if number and int(number[1]) > pages:
            previous.unlink() # Only this script's surplus page render artifacts.
    print('Rendered DOCX to PDF; inspect every rendered page before marking QA passed.')

if __name__ == '__main__':
    parser = argparse.ArgumentParser(); parser.add_argument('--render', action='store_true'); args = parser.parse_args()
    docx = build()
    if args.render: render(docx)
    print(str(docx))
